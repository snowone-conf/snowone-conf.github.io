#!/usr/bin/env python3
"""Создаёт черновики докладов и докладчиков из выгрузки ответов Яндекс Форм.

    python3 tools/import_cfp.py export.xlsx [--only 12,15] [--force] [--no-photos]
    make import-cfp FILE=export.xlsx

Колонки выгрузки сопоставляются полям по tools/cfp_mapping.yaml. Для каждой
строки создаются:
  content/persons/<имя-фамилия>/index.ru.md (+ photo.* по ссылке из формы)
  content/talks/<id-название>/index.ru.md   с draft: true и пустым временем/залом

Существующие файлы не перезаписываются (кроме --force): повторный запуск на той же
выгрузке ничего не меняет. Персона, уже заведённая на сайте (та же папка), не
создаётся заново, а просто указывается в докладе. Поля, которых нет в маппинге
(e-mail, телефон и т.п.), в контент не попадают.

Зависимости: PyYAML; openpyxl — только для .xlsx (для .csv не нужен).
"""
import argparse
import csv
import os
import re
import sys
import urllib.request

import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT = os.path.join(ROOT, "content")
MAPPING = os.path.join(ROOT, "tools", "cfp_mapping.yaml")

TRANSLIT = dict(zip(
    "абвгдеёзийклмнопрстуфхцыэ",
    ["a", "b", "v", "g", "d", "e", "e", "z", "i", "i", "k", "l", "m", "n", "o", "p", "r", "s", "t", "u", "f",
     "kh", "ts", "y", "e"]))
TRANSLIT.update({"ж": "zh", "ч": "ch", "ш": "sh", "щ": "shch", "ю": "yu", "я": "ya", "ъ": "", "ь": ""})
FORMATS = {"доклад": "talk", "talk": "talk", "воркшоп": "workshop", "workshop": "workshop",
           "мастер-класс": "workshop", "обсуждение": "conversation", "дискуссия": "conversation",
           "круглый стол": "conversation"}


def slugify(text, limit=60):
    s = "".join(TRANSLIT.get(c, c) for c in text.lower())
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:limit].rstrip("-") or "item"


def read_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Для .xlsx нужен openpyxl: pip install openpyxl (или выгрузите CSV)")
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h or "").strip() for h in rows[0]]
        return [dict(zip(header, ("" if v is None else str(v) for v in r))) for r in rows[1:] if any(r)]
    with open(path, encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        return [{(k or "").strip(): (v or "") for k, v in r.items()} for r in csv.DictReader(fh, dialect=dialect)]


def pick(row, headers):
    """Значение первой найденной колонки из списка вариантов заголовка."""
    if not headers:
        return ""
    if isinstance(headers, str):
        headers = [headers]
    norm = {k.strip().lower(): v for k, v in row.items()}
    for h in headers:
        v = norm.get(h.strip().lower())
        if v is not None:
            return str(v).strip()
    return ""


def paragraphs(text):
    """Текст из формы -> Markdown: абзацы через пустую строку."""
    text = text.replace("\r\n", "\n").strip()
    return re.sub(r"\n\s*\n+", "\n\n", text)


def write_md(path, front, body, force):
    if os.path.exists(path) and not force:
        return False
    os.makedirs(os.path.dirname(path), exist_ok=True)
    front = {k: v for k, v in front.items() if v not in (None, "", [])}
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("---\n" + yaml.safe_dump(front, allow_unicode=True, sort_keys=False, width=1000) + "---\n")
        if body:
            fh.write(body + "\n")
    return True


def download_photo(url, bundle):
    if not re.match(r"https?://", url):
        return None
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "snowone-import-cfp"})
        with urllib.request.urlopen(req, timeout=30) as r:
            ctype = r.headers.get("Content-Type", "")
            data = r.read()
    except Exception as e:  # сеть, 404, диск Яндекса без прямой ссылки...
        print(f"    фото не скачано ({e}); положите его вручную в {os.path.relpath(bundle, ROOT)}/")
        return None
    ext = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}.get(ctype.split(";")[0].strip())
    if not ext:
        print(f"    по ссылке на фото не картинка ({ctype or 'тип неизвестен'}); положите фото вручную")
        return None
    name = "photo" + ext
    os.makedirs(bundle, exist_ok=True)
    with open(os.path.join(bundle, name), "wb") as fh:
        fh.write(data)
    return name


def contacts(sp, m):
    out = []
    for kind in ("telegram", "github", "website"):
        v = pick(sp, m.get(kind))
        if not v:
            continue
        if kind == "telegram" and not v.startswith("http"):
            v = "https://t.me/" + v.lstrip("@")
        elif kind == "github" and not v.startswith("http"):
            v = "https://github.com/" + v.lstrip("@")
        elif not v.startswith("http"):
            v = "https://" + v
        out.append({"type": kind, "url": v})
    return out


def import_row(row, mapping, season, args):
    rid = pick(row, mapping.get("id")) or "?"
    t = mapping["talk"]
    title = pick(row, t.get("title"))
    if not title:
        print(f"[{rid}] пропущено: нет названия доклада")
        return
    speaker_ids = []
    for m in mapping.get("speakers", []):
        name = pick(row, m.get("name"))
        if not name:
            continue
        pid = slugify(name)
        speaker_ids.append(pid)
        bundle = os.path.join(CONTENT, "persons", pid)
        index = os.path.join(bundle, "index.ru.md")
        if os.path.exists(index) and not args.force:
            print(f"[{rid}] докладчик {name}: уже есть content/persons/{pid}/")
            continue
        photo_url = pick(row, m.get("photo"))
        photo = None if args.no_photos or not photo_url else download_photo(photo_url, bundle)
        write_md(index, {
            "title": name,
            "company": pick(row, m.get("company")),
            "position": pick(row, m.get("position")),
            "photo": photo,
            "weight": 100,
            "contacts": contacts(row, m),
            "draft": True,
        }, paragraphs(pick(row, m.get("bio"))), args.force)
        print(f"[{rid}] докладчик {name}: создан content/persons/{pid}/")

    tid = slugify(f"{rid}-{title}") if rid != "?" else slugify(title)
    lang = pick(row, t.get("language")).lower()
    created = write_md(os.path.join(CONTENT, "talks", tid, "index.ru.md"), {
        "title": title,
        "format": FORMATS.get(pick(row, t.get("format")).lower(), "talk"),
        "speakers": speaker_ids,
        "language": "en" if lang.startswith(("en", "англ")) else "ru",
        "summary": pick(row, t.get("summary")),
        "season": season,
        "cfp_id": rid,
        "draft": True,
    }, paragraphs(pick(row, t.get("description"))), args.force)
    print(f"[{rid}] доклад «{title}»: " + ("создан" if created else "уже есть") + f" content/talks/{tid}/")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("file", help="выгрузка ответов (.xlsx или .csv)")
    ap.add_argument("--only", help="номера ответов через запятую (колонка id)")
    ap.add_argument("--force", action="store_true", help="перезаписать существующие файлы")
    ap.add_argument("--no-photos", action="store_true", help="не скачивать фото")
    ap.add_argument("--season", type=int, help="сезон докладов (по умолчанию — из data/conference.yaml)")
    ap.add_argument("--mapping", default=MAPPING)
    args = ap.parse_args()

    with open(args.mapping, encoding="utf-8") as fh:
        mapping = yaml.safe_load(fh)
    with open(os.path.join(ROOT, "data", "conference.yaml"), encoding="utf-8") as fh:
        season = args.season or yaml.safe_load(fh)["season"]
    rows = read_rows(args.file)
    if args.only:
        wanted = {x.strip() for x in args.only.split(",")}
        rows = [r for r in rows if pick(r, mapping.get("id")) in wanted]
    if not rows:
        sys.exit("в выгрузке нет подходящих строк")
    for row in rows:
        import_row(row, mapping, season, args)
    print("\nЧерновики созданы с draft: true. Заполните день, время и зал, уберите draft и проверьте: make check")


if __name__ == "__main__":
    main()
